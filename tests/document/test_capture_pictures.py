"""
Les captures prises vont dans le document, à la place de leur emplacement.

Une image déposée dans le dossier des captures, sous le nom que `--captures`
lui donne, est insérée là où le plan la désigne : plus de repère « 🖼 » à
remplacer à la main. Ce qui n'a pas de capture garde son emplacement. Et une
régénération remplace l'image par la nouvelle, sans la dupliquer.
"""

import os
import shutil
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn

from main import Options, generate
from src.core import console
from src.core.config import DEFAULT_CAPTURES_DIR, DEFAULT_CONFIG_PATH
from src.gui_automator.fake import solid_png
from src.gui_automator.library import CaptureLibrary

FIXTURE = os.path.join(os.path.dirname(__file__), "..", "fixtures", "rapport_test")
PAGE = "page_ventes"
PAGE_TITLE = "Synthèse commerciale"


class CapturePicturesTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._temp = tempfile.mkdtemp()
        cls.project = os.path.join(cls._temp, "rapport_test")
        shutil.copytree(FIXTURE, cls.project)
        cls.library = CaptureLibrary(os.path.join(cls.project, DEFAULT_CAPTURES_DIR))
        # Une page entière, bien plus large que le texte, et un visuel étroit.
        cls.library.write(PAGE, "_page", solid_png(1874, 1104, (200, 30, 30)))
        cls.library.write(PAGE, "v_evolution", solid_png(150, 90, (30, 30, 200)))
        cls._generate()

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls._temp, ignore_errors=True)

    @classmethod
    def _generate(cls):
        options = Options(
            pbip_path=os.path.join(cls.project, "Rapport.pbip"),
            config_path=DEFAULT_CONFIG_PATH,
            interactive=False,
            pause=False,
        )
        with console.silenced():
            cls.output_dir = generate(options)
        cls.document = Document(os.path.join(cls.output_dir, "documentation_Rapport.docx"))

    def _pictures(self) -> list:
        body = self.document.element.body
        return list(body.iter(qn("wp:inline")))

    def test_les_captures_sont_inserees(self):
        self.assertEqual(len(self._pictures()), 2)

    def test_leur_emplacement_disparait(self):
        texts = [p.text for p in self.document.paragraphs]
        self.assertFalse(any(f"page « {PAGE_TITLE} »" in t and "🖼" in t for t in texts))

    def test_sans_capture_l_emplacement_reste(self):
        self.assertTrue(any("🖼" in p.text for p in self.document.paragraphs))

    def test_une_capture_ne_deborde_pas_du_texte(self):
        section = self.document.sections[-1]
        usable = section.page_width - section.left_margin - section.right_margin
        widths = [int(extent.get("cx")) for extent in self._extents()]
        self.assertTrue(all(width <= usable for width in widths))

    def test_une_petite_capture_garde_sa_taille_d_ecran(self):
        # 150 px à 96 px par pouce : 3,97 cm, soit 1 428 750 EMU.
        widths = sorted(int(extent.get("cx")) for extent in self._extents())
        self.assertAlmostEqual(widths[0], 1428750, delta=400)

    def test_regenerer_remplace_l_image_sans_la_dupliquer(self):
        type(self)._generate()
        self.assertEqual(len(self._pictures()), 2)

    def _extents(self):
        return [inline.find(qn("wp:extent")) for inline in self._pictures()]


class PicturesAfterPlaceholdersTest(unittest.TestCase):
    """
    Le cas courant : un document déjà généré avec ses emplacements, puis les
    captures prises. À la régénération, chaque emplacement resté tel quel
    cède la place à son image — sans que l'emplacement demeure à côté.
    """

    def test_l_emplacement_non_touche_cede_la_place_a_l_image(self):
        temp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, temp, ignore_errors=True)
        project = os.path.join(temp, "rapport_test")
        shutil.copytree(FIXTURE, project)
        options = Options(
            pbip_path=os.path.join(project, "Rapport.pbip"),
            config_path=DEFAULT_CONFIG_PATH,
            interactive=False,
            pause=False,
        )
        with console.silenced():
            generate(options)
        library = CaptureLibrary(os.path.join(project, DEFAULT_CAPTURES_DIR))
        library.write(PAGE, "_page", solid_png(400, 240, (200, 30, 30)))
        with console.silenced():
            output = generate(options)

        document = Document(os.path.join(output, "documentation_Rapport.docx"))
        texts = [p.text for p in document.paragraphs]
        self.assertEqual(len(list(document.element.body.iter(qn("wp:inline")))), 1)
        self.assertFalse(any(f"page « {PAGE_TITLE} »" in t and "🖼" in t for t in texts))


if __name__ == "__main__":
    unittest.main()
