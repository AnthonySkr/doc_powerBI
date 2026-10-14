"""
Cartouche de la page de garde : écrit une fois, puis à l'utilisateur.

Il est posé à la création — la date du jour remplie, le reste vide — et la
régénération le reprend tel que l'utilisateur l'a rempli. La page de garde
précède la première ancre : sans cette reprise, le cartouche rempli partirait
en annexe à chaque génération.
"""

import os
import tempfile
import unittest
from datetime import datetime

from docx import Document
from docx.oxml.ns import qn

from src.core import console
from src.core.config import DocConfig
from src.report_generator.word import cover, generate_word_documentation
from tests.document.test_merge_cycle import PLAN, context

_LABELS = ["Version", "Date de mise à jour", "Réalisée par"]
_COVER = {
    "placeholder": "Titre du document",
    "text": "Documentation",
    "table": {
        "date_format": "%d/%m/%Y",
        "columns": [
            {"label": "Version"},
            {"label": "Date de mise à jour", "value": "{{ today }}"},
            {"label": "Réalisée par"},
        ],
    },
}


class CoverTableTest(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.path = os.path.join(directory.name, "doc.docx")
        self.template = os.path.join(directory.name, "template.docx")
        template = Document()
        template.add_paragraph("Titre du document")
        template.add_paragraph("Suite du template")
        template.save(self.template)

    def generate(self, table: dict | None = None):
        options = {**_COVER, "table": _COVER["table"] if table is None else table}
        raw = {**PLAN, "document": {"template": self.template, "cover": options}}
        with console.silenced():
            generate_word_documentation(DocConfig(raw), context(marge="SUM(T[a])"), self.path)

    def cells(self) -> list[list[str]]:
        table = cover.find(Document(self.path))
        self.assertIsNotNone(table, "aucun cartouche dans le document")
        return [
            [cell.xpath("string(.)") for cell in row.iter(qn("w:tc"))]
            for row in table.iter(qn("w:tr"))
        ]

    def fill(self, *values: str, extra_row: str = "") -> None:
        """Remplit le cartouche dans Word : la ligne prévue, et une de plus."""
        document = Document(self.path)
        table = next(t for t in document.tables if t._tbl is cover.find(document))
        for cell, value in zip(table.rows[1].cells, values, strict=False):
            if value:
                cell.text = value
        if extra_row:
            table.add_row().cells[0].text = extra_row
        document.save(self.path)

    # ── Création ──────────────────────────────────────────────────
    def test_cartouche_pose_a_la_creation(self):
        self.generate()

        today = datetime.now().strftime("%d/%m/%Y")  # noqa: DTZ005
        self.assertEqual(self.cells(), [_LABELS, ["", today, ""]])

    def test_cartouche_dans_la_page_de_garde(self):
        """Devant le titre : il reste dans la section de la page de garde."""
        self.generate()

        document = Document(self.path)
        table = cover.find(document)
        title = table.getnext()
        self.assertEqual(title.xpath("string(.)"), "Documentation")

    def test_cartouche_en_bas_de_page(self):
        self.generate()

        position = cover.find(Document(self.path)).xpath("./w:tblPr/w:tblpPr")[0]
        self.assertEqual(position.get(qn("w:tblpYSpec")), "bottom")

    def test_pas_de_cartouche_sans_colonnes(self):
        self.generate(table={})

        self.assertIsNone(cover.find(Document(self.path)))

    def test_cartouche_desactive(self):
        self.generate(table={**_COVER["table"], "enabled": False})

        self.assertIsNone(cover.find(Document(self.path)))

    # ── Régénération ──────────────────────────────────────────────
    def test_cartouche_rempli_repris_tel_quel(self):
        self.generate()
        self.fill("v1.2", "", "A. Dupont", extra_row="v1.3")
        expected = self.cells()

        self.generate()

        self.assertEqual(self.cells(), expected)
        self.assertEqual(expected[1][0], "v1.2")
        self.assertEqual(expected[2][0], "v1.3")

    def test_cartouche_rempli_ne_part_pas_en_annexe(self):
        self.generate()
        self.fill("v1.2", "", "A. Dupont")

        self.generate()

        texts = [p.text for p in Document(self.path).paragraphs]
        self.assertFalse(any("non replacé" in text for text in texts))

    def test_un_seul_cartouche_apres_regenerations(self):
        self.generate()
        self.fill("v1")
        self.generate()
        self.generate()

        document = Document(self.path)
        found = [
            t for t in document.tables if t._tbl.xpath(f".//w:tblCaption[@w:val='{cover.CAPTION}']")
        ]
        self.assertEqual(len(found), 1)

    def test_cartouche_supprime_regenere_neuf(self):
        self.generate()
        self.fill("v1")
        document = Document(self.path)
        table = cover.find(document)
        table.getparent().remove(table)
        document.save(self.path)

        self.generate()

        self.assertEqual(self.cells()[1][0], "")


if __name__ == "__main__":
    unittest.main()
